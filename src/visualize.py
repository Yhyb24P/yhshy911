"""从冻结结果生成输入、时空场、机理、验证和决策图件；不重算 PDE。"""
import hashlib
import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from metrics import drying_rate_from_flux, profile_moments
from props import DryingRoom, Radius, R0
from solver import map_to_physical, surface_from_state
from utils import atomic_json_dump, input_signature, setup_plot


ROOT = Path(__file__).resolve().parent.parent
TABLES = ROOT / "results" / "tables"
FIGURES = ROOT / "results" / "figures"
DATA = ROOT / "data"


def _save(fig, name, png=False):
    fig.tight_layout()
    pdf = FIGURES / f"{name}.pdf"
    fig.savefig(pdf, bbox_inches="tight")
    files = [pdf.name]
    if png:
        raster = FIGURES / f"{name}.png"
        fig.savefig(raster, dpi=300, bbox_inches="tight")
        files.append(raster.name)
    plt.close(fig)
    return files


def _sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _heatmap(times, coordinates, values, xlabel, ylabel, color_label, name,
             cmap="viridis", contour=None, boundary=None):
    setup_plot()
    fig, ax = plt.subplots(figsize=(7.2, 4.6))
    mesh = ax.pcolormesh(times, coordinates, np.asarray(values).T,
                         shading="auto", cmap=cmap, rasterized=True)
    fig.colorbar(mesh, ax=ax, label=color_label)
    if contour is not None:
        finite = np.nanmin(values), np.nanmax(values)
        if finite[0] <= contour <= finite[1]:
            ax.contour(times, coordinates, np.asarray(values).T,
                       levels=[contour], colors="white", linewidths=1.0)
    if boundary is not None:
        ax.plot(times, boundary, color="black", lw=1.2, label="R(t)")
        ax.legend()
    ax.set_xlabel(xlabel)
    ax.set_ylabel(ylabel)
    return _save(fig, name, png=True)


def _xlsx_series(path, sheet, stop, every=60):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))[1:]
    times, rows = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        t = int(row[0])
        if t > stop:
            break
        if t % every == 0:
            times.append(t)
            rows.append(row[1:])
    wb.close()
    return np.asarray(times), np.asarray(header, dtype=float), np.asarray(rows, dtype=float)


def _prepend_initial(times, values, initial):
    return np.insert(np.asarray(times), 0, 0.0), np.vstack([
        np.full((1, values.shape[1]), initial), values])


def _q4_grids(df):
    fixed_cm = np.arange(0.0, 2.0, 0.1)
    r_grid = np.linspace(0.0, 2.0, 101)
    xi_grid = np.linspace(0.0, 1.0, 101)
    radius = Radius(DATA / "附件2.xlsx")
    physical, material, R_cm = [], [], []
    for _, row in df.iterrows():
        t = float(row["t"])
        R = float(radius(t)) * 100.0
        vals = row[[f"X_{r:.1f}" for r in fixed_cm]].to_numpy(dtype=float)
        valid = np.isfinite(vals) & (fixed_cm <= R + 1e-10)
        r_src = np.append(fixed_cm[valid], R)
        x_src = np.append(vals[valid], float(row["X_surface"]))
        order = np.argsort(r_src)
        r_src, x_src = r_src[order], x_src[order]
        r_src, unique = np.unique(r_src, return_index=True)
        x_src = x_src[unique]
        p = np.interp(np.minimum(r_grid, R), r_src, x_src)
        p[r_grid > R + 1e-10] = np.nan
        physical.append(p)
        material.append(np.interp(xi_grid * R, r_src, x_src))
        R_cm.append(R)
    return r_grid, xi_grid, np.asarray(physical), np.asarray(material), np.asarray(R_cm)


def _signature():
    return input_signature([
        __file__, ROOT / "src" / "metrics.py", DATA / "附件1.xlsx",
        DATA / "附件2.xlsx", TABLES / "q1_samples.csv", TABLES / "q2_60s.csv",
        TABLES / "q2_meta.json", TABLES / "q4_samples.csv", TABLES / "q4_meta.json",
        TABLES / "grid_convergence.json", TABLES / "time_convergence.json",
        TABLES / "q_sens.json", TABLES / "quality_pareto.csv",
        TABLES / "quality_profiles.csv", TABLES / "quality_grid_convergence.json",
    ], settings=("pdf", "heatmap-png", "q4-interpolate-frozen-physical-samples"))


def run():
    FIGURES.mkdir(parents=True, exist_ok=True)
    made = []
    room = DryingRoom(DATA / "附件1.xlsx")

    # F1：实测环境与平台延拓。
    setup_plot()
    fig, ax1 = plt.subplots(figsize=(7.2, 4.2))
    horizon = np.linspace(0, 12 * 3600, 721)
    ax1.plot(room.times / 3600, room.Ta_data, "o", ms=3, label="measured T_a")
    ax1.plot(horizon / 3600, room.Ta(horizon), lw=1.5, label="interpolated/platform T_a")
    ax1.set_xlabel("time / h")
    ax1.set_ylabel("ambient temperature / °C")
    ax2 = ax1.twinx()
    ax2.plot(room.times / 3600, room.Ca_data, "s", ms=3, color="tab:blue", label="measured C_a")
    ax2.plot(horizon / 3600, room.Ca(horizon), color="tab:blue", lw=1.5,
             label="interpolated/platform C_a")
    ax2.set_ylabel("ambient moisture potential / kg kg$^{-1}$")
    lines = ax1.lines + ax2.lines
    ax1.legend(lines, [line.get_label() for line in lines], fontsize=8)
    made += _save(fig, "input_environment")

    # F2/F3：Q1 冻结轨迹。
    q1 = pd.read_csv(TABLES / "q1_samples.csv")
    radii = np.arange(0.0, 2.01, 0.1)
    sample = q1.iloc[::5]
    for prefix, name, label, cmap, initial in (
        ("T", "q1_temperature_heatmap", "T / °C", "magma", 28.0),
        ("X", "q1_moisture_heatmap", "X / kg kg$^{-1}$", "viridis", 2.55),
    ):
        values = sample[[f"{prefix}_{r:.1f}" for r in radii]].to_numpy()
        times, values = _prepend_initial(sample["t"].to_numpy() / 60.0, values, initial)
        made += _heatmap(times, radii, values, "time / min", "radius / cm", label,
                         name, cmap=cmap)
    setup_plot()
    fig, axes = plt.subplots(1, 2, figsize=(9.5, 4.0))
    for seconds in (600, 1200, 1800):
        row = q1[q1["t"] == seconds].iloc[0]
        axes[0].plot(radii, [row[f"T_{r:.1f}"] for r in radii], label=f"{seconds/60:g} min")
        axes[1].plot(radii, [row[f"X_{r:.1f}"] for r in radii], label=f"{seconds/60:g} min")
    axes[0].set(xlabel="radius / cm", ylabel="T / °C")
    axes[1].set(xlabel="radius / cm", ylabel="X / kg kg$^{-1}$")
    for ax in axes:
        ax.legend()
    made += _save(fig, "q1_profiles")

    # F4：Q2 前三小时。
    t2, r2, T2 = _xlsx_series(DATA / "附件3" / "result2.xlsx", "温度", 10800)
    t2, T2 = _prepend_initial(t2 / 3600.0, T2, 28.0)
    q2 = pd.read_csv(TABLES / "q2_60s.csv")
    q2_3h = q2[q2["t"] <= 10800]
    X2 = q2_3h[[f"X_{r:.1f}" for r in radii]].to_numpy()
    tx2, X2 = _prepend_initial(q2_3h["t"].to_numpy() / 3600.0, X2, 2.55)
    made += _heatmap(t2, r2, T2, "time / h", "radius / cm", "T / °C",
                     "q2_temperature_heatmap", cmap="magma")
    made += _heatmap(tx2, radii, X2, "time / h", "radius / cm",
                     "X / kg kg$^{-1}$", "q2_moisture_heatmap")

    # F5/F6：Q3 全程场，以及面积平均含水率–通量干燥速率。
    X3 = q2[[f"X_{r:.1f}" for r in radii]].to_numpy()
    t3, X3 = _prepend_initial(q2["t"].to_numpy() / 3600.0, X3, 2.55)
    made += _heatmap(t3, radii, X3, "time / h", "radius / cm",
                     "X / kg kg$^{-1}$", "q3_moisture_heatmap", contour=0.15)
    Xmean, _ = profile_moments(X3, radii * 1e-2, R0)
    Ca = room.Ca(t3 * 3600.0)
    rate = drying_rate_from_flux(X3[:, -1], Ca, R0)
    setup_plot()
    fig, ax = plt.subplots()
    ax.plot(Xmean, rate * 3600.0)
    ax.set_xlabel("area-weighted mean X / kg kg$^{-1}$")
    ax.set_ylabel("drying rate / h$^{-1}$")
    made += _save(fig, "q3_drying_rate")
    setup_plot()
    fig, ax = plt.subplots()
    ax.plot(t3, X3[:, 0], label="center")
    ax.plot(t3, Xmean, label="area-weighted mean")
    ax.plot(t3, X3[:, -1], label="surface")
    ax.set_xlabel("time / h")
    ax.set_ylabel("X / kg kg$^{-1}$")
    ax.legend()
    made += _save(fig, "q3_moisture_summary")

    with open(TABLES / "q2_meta.json", encoding="utf-8") as f:
        q2_meta = json.load(f)
    q3_event = map_to_physical(np.asarray(q2_meta["X_event"]), radii * 1e-2, R0)
    q3_event[-1] = surface_from_state(
        2, room, q2_meta["t_dry"], np.asarray(q2_meta["T_event"]),
        np.asarray(q2_meta["X_event"]), R0)[1]
    setup_plot()
    fig, ax = plt.subplots()
    for hour in (6, 18, 36, 54):
        row = q2[q2["t"] == hour * 3600].iloc[0]
        ax.plot(radii, [row[f"X_{r:.1f}"] for r in radii], label=f"{hour} h")
    ax.plot(radii, q3_event, lw=2, label=f"event {q2_meta['t_dry']/3600:.2f} h")
    ax.set_xlabel("radius / cm")
    ax.set_ylabel("X / kg kg$^{-1}$")
    ax.legend(fontsize=8)
    made += _save(fig, "q3_profiles")

    # F7/F8：Q4 物理域与材料域双图。
    q4 = pd.read_csv(TABLES / "q4_samples.csv")
    r_grid, xi_grid, q4_phys, q4_mat, R_cm = _q4_grids(q4)
    tq4 = q4["t"].to_numpy() / 3600.0
    made += _heatmap(tq4, r_grid, q4_phys, "time / h", "physical radius / cm",
                     "X / kg kg$^{-1}$", "q4_physical_heatmap", boundary=R_cm)
    made += _heatmap(tq4, xi_grid, q4_mat, "time / h", "material coordinate ξ",
                     "X / kg kg$^{-1}$", "q4_material_heatmap")
    with open(TABLES / "q4_meta.json", encoding="utf-8") as f:
        q4_meta = json.load(f)
    event_R = float(Radius(DATA / "附件2.xlsx")(q4_meta["t_dry"]))
    event_r_cm = np.linspace(0.0, event_R * 100.0, 101)
    q4_event = map_to_physical(np.asarray(q4_meta["X_event"]), event_r_cm * 1e-2, event_R)
    q4_event[-1] = surface_from_state(
        4, room, q4_meta["t_dry"], np.asarray(q4_meta["T_event"]),
        np.asarray(q4_meta["X_event"]), event_R)[1]
    setup_plot()
    fig, ax = plt.subplots()
    for hour in (6, 18, 36, 48):
        k = int(np.argmin(np.abs(tq4 - hour)))
        ax.plot(r_grid, q4_phys[k], label=f"{hour} h")
    ax.plot(event_r_cm, q4_event, lw=2, label=f"event {q4_meta['t_dry']/3600:.2f} h")
    ax.set_xlabel("physical radius / cm")
    ax.set_ylabel("X / kg kg$^{-1}$")
    ax.legend(fontsize=8)
    made += _save(fig, "q4_profiles")

    # F9：收缩尺度与几何扩散倍率。
    radius_model = Radius(DATA / "附件2.xlsx")
    raw_radius = pd.read_excel(DATA / "附件2.xlsx")
    tr = np.linspace(0.0, q4["t"].iloc[-1], 1000)
    R = radius_model(tr)
    setup_plot()
    fig, ax1 = plt.subplots()
    ax1.plot(raw_radius["时间"] / 3600, raw_radius["半径"], "o", ms=2.5,
             label="measured radius")
    ax1.plot(tr / 3600, R * 100, label="PCHIP R(t)")
    ax1.set_xlabel("time / h")
    ax1.set_ylabel("radius / cm")
    ax2 = ax1.twinx()
    ax2.plot(tr / 3600, R / R0, color="tab:green", label="λ=R/R0")
    ax2.plot(tr / 3600, (R0 / R) ** 2, color="tab:red", label="G_R=(R0/R)^2")
    ax2.set_ylabel("dimensionless shrinkage / diffusion scale")
    lines = ax1.lines + ax2.lines
    ax1.legend(lines, [line.get_label() for line in lines], fontsize=8)
    made += _save(fig, "q4_shrinkage_scale")

    # F10：空间与时间收敛。
    with open(TABLES / "grid_convergence.json", encoding="utf-8") as f:
        grid = json.load(f)["t_dry_s"]
    with open(TABLES / "time_convergence.json", encoding="utf-8") as f:
        temporal = json.load(f)["t_dry_s"]
    setup_plot()
    fig, axes = plt.subplots(1, 2, figsize=(9.5, 4.0))
    for q in ("3", "4"):
        Ns = np.array(sorted(map(int, grid[q])))
        axes[0].plot(Ns, [grid[q][str(n)] / 3600 for n in Ns], "o-", label=f"Q{q}")
        dts = np.array(sorted(map(float, temporal[q]), reverse=True))
        axes[1].plot(dts, [temporal[q][str(dt)] / 3600 for dt in dts], "o-", label=f"Q{q}")
    axes[0].set(xlabel="N", ylabel="drying time / h")
    axes[1].set(xlabel="maximum time step / s", ylabel="drying time / h")
    for ax in axes:
        ax.legend()
    made += _save(fig, "convergence")

    # F11：已有 OAT/平台情景的终止时间变化。
    with open(TABLES / "q_sens.json", encoding="utf-8") as f:
        sens = json.load(f)
    base = sens["t_base"]
    items = []
    for group in ("hT", "hm"):
        for key, value in sens[group].items():
            if float(key) != 1.0:
                items.append((f"{group}×{key}", (value - base) / 3600))
    for key, value in sens["platform"].items():
        items.append((key, (value - base) / 3600))
    setup_plot()
    fig, ax = plt.subplots(figsize=(7.2, 5.2))
    labels, changes = zip(*sorted(items, key=lambda item: item[1]))
    ax.barh(labels, changes, color=["tab:blue" if x < 0 else "tab:orange" for x in changes])
    ax.axvline(0, color="black", lw=0.8)
    ax.set_xlabel("change in drying time / h")
    made += _save(fig, "sensitivity")

    # F12/F13：双目标情景与代表方案终点剖面。
    quality = pd.read_csv(TABLES / "quality_pareto.csv")
    profiles = pd.read_csv(TABLES / "quality_profiles.csv")
    with open(TABLES / "quality_grid_convergence.json", encoding="utf-8") as f:
        quality_grid = json.load(f)
    setup_plot()
    fig, ax = plt.subplots()
    hm_span = np.ptp(quality["hm_factor"])
    point_sizes = (np.full(len(quality), 35.0) if hm_span == 0 else
                   25.0 + 55.0 * (quality["hm_factor"] - quality["hm_factor"].min()) / hm_span)
    points = ax.scatter(quality["t_dry_h"], quality["X_std_event"],
                        c=quality["T_platform_C"], cmap="plasma", s=point_sizes)
    front = quality[quality["is_pareto"]].sort_values("t_dry_h")
    ax.plot(front["t_dry_h"], front["X_std_event"], "D-", color="black", lw=1,
            ms=7, mfc="none", label="nondominated set")
    fine = pd.DataFrame(row for row in quality_grid["rows"] if row["N"] == 641)
    ax.scatter(fine["t_dry_s"] / 3600.0, fine["X_std_event"], marker="x", s=55,
               color="tab:green", label="N=641 convergence points")
    baseline = quality[quality["is_baseline"]]
    ax.scatter(baseline["t_dry_h"], baseline["X_std_event"], marker="*", s=130,
               facecolors="none", edgecolors="black", label="measured baseline")
    fig.colorbar(points, ax=ax, label="platform temperature / °C")
    ax.set_xlabel("drying time / h")
    ax.set_ylabel("area-weighted σ_X at event")
    ax.legend()
    made += _save(fig, "pareto_time_uniformity", png=True)

    setup_plot()
    fig, ax = plt.subplots()
    representatives = quality[quality["representative"].fillna("") != ""]
    for _, scenario in representatives.iterrows():
        profile = profiles[profiles["scenario"] == scenario["scenario"]]
        ax.plot(profile["r_cm"], profile["X_event"],
                label=(f"{scenario['representative']} "
                       f"({scenario['T_platform_C']:.3g} °C, h_m×{scenario['hm_factor']:g})"))
    ax.axhline(0.15, color="gray", ls="--", lw=1)
    ax.set_xlabel("radius / cm")
    ax.set_ylabel("X at drying event / kg kg$^{-1}$")
    ax.legend(fontsize=8)
    made += _save(fig, "pareto_profiles", png=True)

    atomic_json_dump({"signature": _signature(), "files": made,
                      "source_sha256": {
                          "quality_pareto.csv": _sha256(TABLES / "quality_pareto.csv"),
                          "quality_profiles.csv": _sha256(TABLES / "quality_profiles.csv"),
                          "quality_grid_convergence.json": _sha256(
                              TABLES / "quality_grid_convergence.json"),
                      },
                      "output_sha256": {name: _sha256(FIGURES / name) for name in made},
                      "reconstruction_notes": {
                          "q1_heatmaps": "frozen 0.1 cm output samples; cannot resolve the initial thin boundary layer",
                          "q4_heatmaps": "interpolation of frozen 0.1 cm physical samples; not the raw 321-cell field",
                      }},
                     TABLES / "visualization_meta.json")
    print(f"[visualize] 写入 {len(made)} 个图件；未运行 PDE")


if __name__ == "__main__":
    run()
