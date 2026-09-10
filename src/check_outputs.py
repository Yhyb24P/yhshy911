"""审计四个正式 Excel 产物是否满足 model.md §17。

本脚本只读文件，不运行 PDE。大工作簿使用 XML 流式扫描时间列，避免把
result2.xlsx 的数百万个单元格全部载入内存。
"""
import json
import math
import os
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import model1
import model2
import model3
import model4
import generate_tables
import validate
from xlsx_io import HEADER, Q4_R_COLS, R_COLS


ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "附件3"
TABLES = ROOT / "results" / "tables"
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _assert(condition, message):
    if not condition:
        raise AssertionError(message)


def _sheet_xml_paths(path):
    """按工作簿顺序返回工作表 XML 路径。"""
    with ZipFile(path) as zf:
        root = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid_to_target = {
        rel.attrib["Id"]: rel.attrib["Target"].lstrip("/") for rel in rels
    }
    paths = []
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    for sheet in root.find(f"{NS}sheets"):
        target = rid_to_target[sheet.attrib[rel_ns]]
        paths.append(target if target.startswith("xl/") else "xl/" + target)
    return paths


def _time_axis(path, sheet_xml):
    """流式读取 A 列，返回数据行数、首末时刻和步长一致性。"""
    count = 0
    first = last = previous = None
    step = None
    regular = True
    with ZipFile(path) as zf, zf.open(sheet_xml) as stream:
        for _event, elem in ET.iterparse(stream, events=("end",)):
            if elem.tag != f"{NS}row":
                continue
            row_number = int(elem.attrib["r"])
            if row_number >= 2:
                cell = elem.find(f"{NS}c")
                _assert(cell is not None and cell.attrib.get("r", "").startswith("A"),
                        f"{path.name}: 第 {row_number} 行缺少时间")
                value = cell.find(f"{NS}v")
                _assert(value is not None, f"{path.name}: 第 {row_number} 行时间为空")
                t = float(value.text)
                if first is None:
                    first = t
                if previous is not None:
                    delta = t - previous
                    if step is None:
                        step = delta
                    regular &= abs(delta - step) < 1e-9
                previous = last = t
                count += 1
            elem.clear()
    return count, first, last, step, regular


def _check_layout(path, sheet_names, headers):
    wb = load_workbook(path, read_only=True, data_only=False)
    _assert(wb.sheetnames == sheet_names,
            f"{path.name}: 工作表应为 {sheet_names}，实际为 {wb.sheetnames}")
    for ws in wb.worksheets:
        row1 = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        _assert(list(row1) == [HEADER] + list(headers), f"{path.name}/{ws.title}: 表头错误")
        row2 = next(ws.iter_rows(min_row=2, max_row=2, values_only=False))
        _assert(row2[0].number_format == "0", f"{path.name}/{ws.title}: 时间格式应为整数")
        _assert(all(c.value is None or c.number_format == "0.0000" for c in row2[1:]),
                f"{path.name}/{ws.title}: 数值格式应为四位小数")
    wb.close()


def _check_book(name, sheet_names, headers, expected_rows, first, last, step):
    path = DATA / name
    _assert(path.exists(), f"缺少 {path}")
    _check_layout(path, sheet_names, headers)
    axes = [_time_axis(path, xml) for xml in _sheet_xml_paths(path)]
    for got in axes:
        n, t0, t1, dt, regular = got
        _assert(n == expected_rows, f"{name}: 数据行应为 {expected_rows}，实际为 {n}")
        _assert(t0 == first and t1 == last, f"{name}: 时间范围应为 {first}..{last}，实际为 {t0}..{t1}")
        if expected_rows > 1:
            _assert(regular and dt == step, f"{name}: 时间步长并非恒定 {step} s")
    _assert(len(set(axes)) == 1, f"{name}: 各工作表时间轴不一致")
    print(f"[通过] {name}: {expected_rows} 行，t={first:g}..{last:g} s，步长 {step:g} s")


def _load_meta(name):
    path = TABLES / name
    _assert(path.exists(), f"缺少或尚未完成 {path}")
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _check_event(meta, label):
    _assert(meta["t_dry"] > 0, f"{label}: 终止时间无效")
    xmax = max(meta["X_event"])
    _assert(abs(xmax - 0.15) <= 2e-8,
            f"{label}: 事件场 max(X)={xmax:.12g} 未定位到 0.15")


def _selected_rows(path, sheet, times):
    """顺序抽取指定时刻，供数值内容核对。"""
    wanted = {int(t) for t in times}
    out = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    for row in ws.iter_rows(min_row=2, values_only=True):
        t = int(row[0])
        if t in wanted:
            out[t] = dict(zip(header[1:], row[1:]))
        if t >= max(wanted):
            break
    wb.close()
    _assert(out.keys() == wanted, f"{path.name}/{sheet}: 数值抽查时刻不完整")
    return out


def _assert_value(got, expected, label):
    """Excel 四位小数应等于轨迹文件按相同规则舍入后的值。"""
    if pd.isna(expected):
        _assert(got is None, f"{label}: 越材位置应为空，实际为 {got}")
    else:
        _assert(got is not None and abs(float(got) - round(float(expected), 4)) < 5e-9,
                f"{label}: 实际 {got}，轨迹舍入值 {round(float(expected), 4)}")


def _content_spot_checks(q2, q4):
    """用独立轨迹/事件元数据抽查正式工作簿和论文表的数值内容。"""
    q1 = pd.read_csv(TABLES / "q1_samples.csv").set_index("t")
    q2_csv = pd.read_csv(TABLES / "q2_60s.csv").set_index("t")
    q4_csv = pd.read_csv(TABLES / "q4_samples.csv").set_index("t")

    q1_times = (1, 100, 1800)
    for sheet, prefix in (("温度", "T"), ("水分浓度", "X")):
        got = _selected_rows(DATA / "result1.xlsx", sheet, q1_times)
        for t in q1_times:
            for r in (0.0, 1.0, 2.0):
                _assert_value(got[t][r], q1.loc[t, f"{prefix}_{r}"],
                              f"result1/{sheet}, t={t}, r={r}")

    q3_times = (60, 10800, int(q2["t_dry"] // 60) * 60)
    for book in ("result2.xlsx", "result3.xlsx"):
        sheet = "水分浓度" if book == "result2.xlsx" else "Sheet1"
        got = _selected_rows(DATA / book, sheet, q3_times)
        for t in q3_times:
            for r in (0.0, 1.0, 2.0):
                _assert_value(got[t][r], q2_csv.loc[t, f"X_{r}"],
                              f"{book}, t={t}, r={r}")

    q4_times = (60, 10800, int(q4["t_dry"] // 60) * 60)
    got = _selected_rows(DATA / "result4.xlsx", "Sheet1", q4_times)
    for t in q4_times:
        for header, column in ((0.0, "X_0.0"), (1.0, "X_1.0"),
                               (1.5, "X_1.5"), ("药材表面", "X_surface")):
            _assert_value(got[t][header], q4_csv.loc[t, column],
                          f"result4, t={t}, r={header}")

    tables_meta = _load_meta("paper_tables_meta.json")
    _assert(tables_meta["signature"] == generate_tables._signature(),
            "表1–6 生成结果指纹已过期")
    _assert(tables_meta["q3_event_s"] == q2["t_dry"], "表5 精确事件时刻错误")
    _assert(tables_meta["q4_event_s"] == q4["t_dry"], "表6 精确事件时刻错误")
    room = generate_tables.DryingRoom(ROOT / "data" / "附件1.xlsx")
    event_rows = {
        5: generate_tables._event_values(2, q2, generate_tables.RADII, room),
        6: generate_tables._event_values(
            4, q4, generate_tables.Q4_FIXED, room,
            generate_tables.Radius(ROOT / "data" / "附件2.xlsx"), add_surface=True),
    }
    for number, expected_rows in {1: 7, 2: 7, 3: 7, 4: 7, 5: 11, 6: 10}.items():
        path = TABLES / f"paper_table{number}.csv"
        _assert(path.exists(), f"缺少 {path}")
        df = pd.read_csv(path)
        _assert(len(df) == expected_rows, f"表{number}: 应有 {expected_rows} 行")
        _assert(np.isfinite(pd.to_numeric(df.iloc[-1, 1:], errors="coerce")).any(),
                f"表{number}: 末行没有有效数值")
        if number in (5, 6):
            _assert(str(df.iloc[-1, 0]).startswith("烘干结束时间（"),
                    f"表{number}: 缺少精确终止事件行")
            vmax = pd.to_numeric(df.iloc[-1, 1:], errors="coerce").max()
            _assert(abs(vmax - 0.15) <= 5.1e-5,
                    f"表{number}: 终止行 max(X)={vmax} 不符合 0.15")
            actual = pd.to_numeric(df.iloc[-1, 1:], errors="coerce").to_numpy()
            expected = np.round(event_rows[number], 4)
            _assert(np.allclose(actual, expected, atol=5e-9, equal_nan=True),
                    f"表{number}: 终止行场值与事件元数据不一致")
    print("[通过] 数值内容抽查：正式工作簿↔轨迹文件，表5/6↔连续事件元数据")


def run():
    q1 = _load_meta("q1_meta.json")
    q2 = _load_meta("q2_meta.json")
    q3 = _load_meta("q3_meta.json")
    q4 = _load_meta("q4_meta.json")
    _assert(q1.get("N") == 2561, f"Q1 正式结果应使用 N=2561，实际 N={q1.get('N')}")
    _assert(q2.get("N") == 321, f"Q2/Q3 正式结果应使用 N=321，实际 N={q2.get('N')}")
    _assert(q3.get("N") == 321, f"Q3 正式结果应使用 N=321，实际 N={q3.get('N')}")
    _assert(q4.get("N") == 321, f"Q4 正式结果应使用 N=321，实际 N={q4.get('N')}")
    _assert(q1.get("signature") == model1._signature(2561), "Q1 缓存指纹已过期")
    _assert(q2.get("signature") == model2._signature(321), "Q2 缓存指纹已过期")
    _assert(q3.get("signature") == model3._signature(321), "Q3 缓存指纹已过期")
    _assert(q4.get("signature") == model4._signature(321), "Q4 缓存指纹已过期")
    _assert(q3.get("t_dry") == q2.get("t_dry"), "Q3 与 Q2 终止事件元数据不一致")
    _check_event(q2, "Q3")
    _check_event(q4, "Q4")

    q3_rows = math.floor(q2["t_dry"] / 60.0)
    q4_rows = math.floor(q4["t_dry"] / 60.0)
    _check_book("result1.xlsx", ["温度", "水分浓度"], R_COLS,
                1800, 1, 1800, 1)
    _check_book("result2.xlsx", ["温度", "水分浓度"], R_COLS,
                q2["n_rows"], 1, q2["n_rows"], 1)
    _check_book("result3.xlsx", ["Sheet1"], R_COLS,
                q3_rows, 60, q3_rows * 60, 60)
    _check_book("result4.xlsx", ["Sheet1"], Q4_R_COLS + ["药材表面"],
                q4_rows, 60, q4_rows * 60, 60)
    grid = _load_meta("grid_convergence.json")
    grid_ns = tuple(sorted(int(n) for n in grid["t_dry_s"]["3"]))
    _assert(grid.get("signature") == validate._signature(("grid", grid_ns)),
            "网格收敛结果指纹已过期")
    grid_q3_n321 = grid["t_dry_s"]["3"]["321"]
    _assert(abs(grid_q3_n321 - q2["t_dry"]) < 0.01,
            f"Q3 正式事件与网格验证 N=321 不一致：{q2['t_dry']:.4f} vs {grid_q3_n321:.4f} s")
    _content_spot_checks(q2, q4)
    print(f"[通过] 连续事件：Q3={q2['t_dry']:.4f} s，Q4={q4['t_dry']:.4f} s")


if __name__ == "__main__":
    run()
