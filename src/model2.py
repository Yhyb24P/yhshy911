"""问题 2：整个固定尺寸烘干过程。

附录 3 变物性，参数型双向耦合 T↔X（model.md §8），从 t=0 独立求解，
计算持续到固定半径烘干终止时刻 t_dry^(3)（model.md §8、§17.2）。

输出节奏（model.md §17.2「时间首行 t=1 s」）：result2.xlsx 按 1 s 逐行，
首行 t=1 s、不写 t=0，覆盖整个烘干区间（t_dry^(3)≈2.07e5 s，约 2e5 行，
远低于 Excel 单表 1,048,576 行上限）；论文表 3/4 只截取前 3 h。
另写严格 60 s 间隔 CSV（q2_60s.csv）供 model3/result3 复用；精确事件保存在元数据中。
"""
import json
import os
import time as _time

import numpy as np
from openpyxl import Workbook

from props import DryingRoom
from solver import FVMSolver
from utils import atomic_json_dump, input_signature, staged_path
from xlsx_io import R_COLS, data_row, header_row

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
META = os.path.join(ROOT, "results/tables/q2_meta.json")
CSV = os.path.join(ROOT, "results/tables/q2_60s.csv")
INTERVAL = 1.0      # result2.xlsx 输出间隔（model.md §17.2 首行 t=1 s）
CSV_EVERY = 60.0    # q2_60s.csv 采样间隔（供 model3/result3，model.md §17.3）


def _signature(N):
    return input_signature([
        __file__, os.path.join(ROOT, "src/solver.py"),
        os.path.join(ROOT, "src/props.py"), os.path.join(ROOT, "src/xlsx_io.py"),
        os.path.join(ROOT, "data/附件1.xlsx"),
    ], settings=(N, INTERVAL, CSV_EVERY))


def run(N=321, reuse=True):
    out = os.path.join(ROOT, "data/附件3/result2.xlsx")
    signature = _signature(N)
    if reuse and all(os.path.exists(p) for p in (META, CSV, out)):
        meta = json.load(open(META))
        if meta.get("N") == N and meta.get("signature") == signature:
            print(f"[model2] 复用已有结果：t_dry^(3)={meta['t_dry']:.2f} s（{meta['t_dry']/3600:.1f} h）")
            return meta

    # 一旦进入重算便使旧元数据失效；正式产物在临时文件完整写好后再替换。
    if os.path.exists(META):
        os.remove(META)
    tmp_out, tmp_csv = staged_path(out), staged_path(CSV)

    t_start = _time.time()
    room = DryingRoom(os.path.join(ROOT, "data/附件1.xlsx"))
    s = FVMSolver(2, room, N=N)
    r_phys = np.array(R_COLS) * 1e-2

    wb = Workbook(write_only=True)
    ws_T = wb.create_sheet("温度")
    ws_X = wb.create_sheet("水分浓度")
    ws_T.append(header_row(ws_T, R_COLS))
    ws_X.append(header_row(ws_X, R_COLS))

    os.makedirs(os.path.dirname(CSV), exist_ok=True)
    csv_f = open(tmp_csv, "w")
    csv_f.write("t," + ",".join(f"X_{c}" for c in R_COLS) + "\n")

    n_rows = 0
    k = 0.0
    while True:
        k += INTERVAL
        lo_solver = s.clone()
        s.advance_to(k)
        if s.g() <= 0.0:
            t_dry, T_ev, X_ev = s.refine_event(lo_solver, k)
            break
        T, X, _ = s.sample(r_phys)
        ws_T.append(data_row(ws_T, int(k), T))
        ws_X.append(data_row(ws_X, int(k), X))
        n_rows += 1
        if int(k) % int(CSV_EVERY) == 0:             # 60 s 采样行（供 model3）
            csv_f.write(f"{int(k)}," + ",".join(f"{v:.6f}" for v in X) + "\n")
        if int(k) % 3600 == 0:
            print(f"    t={int(k):>10d} s   X_max={s.X.max():.4f}", flush=True)
    wb.save(tmp_out)
    print(f"[model2] 写入 {out}（{n_rows} 行 × {len(R_COLS)} 列 × 2 工作表，1 s 间隔）")

    csv_f.close()
    os.replace(tmp_out, out)
    os.replace(tmp_csv, CSV)

    meta = dict(N=N, signature=signature, t_dry=t_dry, n_rows=n_rows,
                T_event=T_ev.tolist(), X_event=X_ev.tolist())
    atomic_json_dump(meta, META)
    print(f"[model2] t_dry^(3) = {t_dry:.2f} s = {t_dry/3600:.2f} h")
    print(f"[model2] 完成，耗时 {_time.time() - t_start:.1f} s")
    return meta


if __name__ == "__main__":
    run()
