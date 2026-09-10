"""从正式结果自动提取题面表 1–6。

规则时刻直接取已冻结的 CSV/Excel 数值；表 5、表 6 的最后一行由事件
元数据中的 cell 场映射得到，因此使用连续终止时刻，而不是最后一个 60 s
采样点。输出位于 results/tables/paper_table[1-6].csv，统一保留四位小数。
"""
import csv
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from props import DryingRoom, Radius, R0, T0, X0
from solver import map_to_physical, surface_from_state
from utils import atomic_json_dump, input_signature, staged_path


ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "附件3"
OUT = ROOT / "results" / "tables"
RADII = [0.0, 0.5, 1.0, 1.5, 2.0]
Q4_FIXED = [0.0, 0.5, 1.0, 1.5]


def _load_json(name):
    with open(OUT / name, encoding="utf-8") as f:
        return json.load(f)


def _xlsx_rows(path, sheet, times):
    """顺序读取大工作簿，只保留指定时刻的行。"""
    wanted = {int(t) for t in times}
    rows = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    for row in ws.iter_rows(min_row=2, values_only=True):
        t = int(row[0])
        if t in wanted:
            rows[t] = dict(zip(header[1:], row[1:]))
        if t >= max(wanted):
            break
    wb.close()
    missing = wanted - rows.keys()
    if missing:
        raise ValueError(f"{path.name}/{sheet} 缺少时刻 {sorted(missing)}")
    return rows


def _write_table(number, time_header, position_headers, rows):
    path = OUT / f"paper_table{number}.csv"
    tmp = staged_path(path)
    try:
        with open(tmp, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, lineterminator="\n")
            writer.writerow([time_header] + position_headers)
            for label, values in rows:
                writer.writerow([label] + ["" if np.isnan(v) else f"{v:.4f}"
                                           for v in np.asarray(values, dtype=float)])
        Path(tmp).replace(path)
    finally:
        Path(tmp).unlink(missing_ok=True)
    print(f"[tables] 写入 {path}")


def _event_values(q, meta, radii_cm, room, radius=None, add_surface=False):
    t = float(meta["t_dry"])
    T = np.asarray(meta["T_event"], dtype=float)
    X = np.asarray(meta["X_event"], dtype=float)
    R = R0 if radius is None else float(radius(t))
    values = map_to_physical(X, np.asarray(radii_cm) * 1e-2, R)
    surface = surface_from_state(q, room, t, T, X, R)[1]
    if add_surface:
        values = np.append(values, surface)
    elif abs(radii_cm[-1] * 1e-2 - R) < 1e-10:
        values[-1] = surface
    return values


def _signature():
    return input_signature([
        __file__, OUT / "q1_samples.csv", DATA / "result2.xlsx",
        OUT / "q2_60s.csv", OUT / "q2_meta.json",
        OUT / "q4_samples.csv", OUT / "q4_meta.json",
    ], settings=(tuple(RADII), tuple(Q4_FIXED), "four-decimal"))


def run():
    OUT.mkdir(parents=True, exist_ok=True)
    q1 = pd.read_csv(OUT / "q1_samples.csv").set_index("t")
    q1_times = [100, 300, 600, 900, 1200, 1500, 1800]
    radius_headers = [f"{r:g} cm" for r in RADII]
    _write_table(1, "时间/s", radius_headers,
                 [(t, [q1.loc[t, f"T_{r}"] for r in RADII]) for t in q1_times])
    _write_table(2, "时间/s", radius_headers,
                 [(t, [q1.loc[t, f"X_{r}"] for r in RADII]) for t in q1_times])

    q2_times = [1800, 3600, 5400, 7200, 9000, 10800]
    Trows = _xlsx_rows(DATA / "result2.xlsx", "温度", q2_times)
    Xrows = _xlsx_rows(DATA / "result2.xlsx", "水分浓度", q2_times)
    table3 = [(0, [T0] * len(RADII))]
    table4 = [(0, [X0] * len(RADII))]
    for t in q2_times:
        table3.append((t / 3600, [Trows[t][r] for r in RADII]))
        table4.append((t / 3600, [Xrows[t][r] for r in RADII]))
    _write_table(3, "时间/h", radius_headers, table3)
    _write_table(4, "时间/h", radius_headers, table4)

    room = DryingRoom(ROOT / "data" / "附件1.xlsx")
    q2_meta = _load_json("q2_meta.json")
    q2_60 = pd.read_csv(OUT / "q2_60s.csv").set_index("t")
    q3_hours = list(range(0, int(q2_meta["t_dry"] // 21600) * 6 + 1, 6))
    table5 = [(0, [X0] * len(RADII))]
    for hour in q3_hours[1:]:
        t = hour * 3600
        table5.append((hour, [q2_60.loc[t, f"X_{r}"] for r in RADII]))
    q3_event = _event_values(2, q2_meta, RADII, room)
    q3_label = f"烘干结束时间（{q2_meta['t_dry']/3600:.4f} h）"
    table5.append((q3_label, q3_event))
    _write_table(5, "时间/h", radius_headers, table5)

    q4_meta = _load_json("q4_meta.json")
    q4_60 = pd.read_csv(OUT / "q4_samples.csv").set_index("t")
    q4_hours = list(range(0, int(q4_meta["t_dry"] // 21600) * 6 + 1, 6))
    q4_headers = [f"{r:g} cm" for r in Q4_FIXED] + ["药材表面"]
    table6 = [(0, [X0] * (len(Q4_FIXED) + 1))]
    for hour in q4_hours[1:]:
        row = q4_60.loc[hour * 3600]
        table6.append((hour, [row[f"X_{r}"] for r in Q4_FIXED] + [row["X_surface"]]))
    radius = Radius(ROOT / "data" / "附件2.xlsx")
    q4_event = _event_values(4, q4_meta, Q4_FIXED, room, radius, add_surface=True)
    q4_label = f"烘干结束时间（{q4_meta['t_dry']/3600:.4f} h）"
    table6.append((q4_label, q4_event))
    _write_table(6, "时间/h", q4_headers, table6)

    meta = {
        "signature": _signature(),
        "q3_event_s": q2_meta["t_dry"],
        "q4_event_s": q4_meta["t_dry"],
        "rows": {str(i): sum(1 for _ in open(OUT / f"paper_table{i}.csv",
                                               encoding="utf-8-sig")) - 1
                 for i in range(1, 7)},
    }
    atomic_json_dump(meta, OUT / "paper_tables_meta.json")
    print(f"[tables] 精确事件行：Q3={q2_meta['t_dry']:.4f} s，Q4={q4_meta['t_dry']:.4f} s")
    return meta


if __name__ == "__main__":
    run()
